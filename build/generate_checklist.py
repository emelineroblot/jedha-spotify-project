"""Remplit la Compliance Checklist officielle (contexte/Compliance_Checklist_Spotify.xlsx)
et l'enregistre dans livrables/02-compliance-checklist.xlsx."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "contexte" / "Compliance_Checklist_Spotify.xlsx"
OUT = ROOT / "livrables" / "02-compliance-checklist.xlsx"

# (statut actuel estimé, owner, action plan, échéance) par ligne, dans l'ordre du template
ROWS = [
    ("Partiel", "DPO", "Registre des traitements (art. 30) et base légale documentée par finalité (contrat / intérêt légitime / consentement / obligation légale). Data mapping OneTrust. Privacy notice réécrite par finalité, avec durées de conservation et transferts (point sanctionné par l'IMY en 2023).", "M4"),
    ("Partiel", "DPO + Privacy Team + Engineering", "Couvrir les 7 droits (accès complet avec source/destinataires/durées/transferts, rectification, effacement, portabilité, opposition, limitation, décision automatisée). Pipeline d'effacement automatisé < 30 j sur tous les systèmes du domaine. Privacy Team (2 analystes) avec SLA 1 mois.", "M6 (User Data), M12 (tous domaines)"),
    ("Partiel", "DPO + Product", "Plateforme de gestion des consentements (OneTrust) granulaire par finalité, aussi simple à retirer qu'à donner. Consentement explicite pour publicité ciblée et partage à des tiers uniquement ; contrat / intérêt légitime pour le service. Consentement parental pour les mineurs (art. 8).", "M2 (déploiement), M6 (audit)"),
    ("Partiel", "DPO + Head of Engineering", "Protocole d'incident : détection SIEM Splunk, qualification sous 24 h, notification autorité sous 72 h, information des personnes si risque élevé. Runbook + exercice annuel. Registre des violations.", "M5 (runbook), M12 (exercice)"),
    ("Oui", "CEO / Board", "DPO confirmé, rattaché à la direction générale hors ligne hiérarchique du CDO (indépendance, art. 38), enregistré auprès de l'autorité, dirige la Privacy Team, membre du Data Governance Committee.", "M1"),
    ("Partiel", "Steward Ads + Product + Legal", "Lien « Do Not Sell or Share My Personal Information » visible (site et app), prise en compte du signal Global Privacy Control, opt-out propagé aux partenaires publicitaires sous 15 jours, suivi OneTrust.", "M9 (domaine Ads)"),
    ("Partiel", "Privacy Team", "Portail libre-service (accès, suppression, correction CPRA) + traitement manuel par la Privacy Team, réponse < 45 jours, vérification d'identité, journal des demandes.", "M6"),
    ("Oui", "Product Managers", "Contrôle produit : aucune dégradation de service, de prix ou de qualité après exercice d'un droit ; test de non-régression à chaque release ; clause explicite dans la privacy notice.", "M6 (vérification)"),
    ("Oui (périmètre délégué)", "Head of Engineering + Steward Payments", "Paiements traités par un prestataire certifié PCI-DSS, aucun numéro de carte stocké chez Spotify (tokenisation). Périmètre résiduel : page de paiement et redirections sécurisées, segmentation réseau, questionnaire SAQ-A annuel.", "M10 (domaine Payments)"),
    ("Oui (périmètre délégué)", "Steward Payments", "Jetons et données de facturation classés « Sensible » : chiffrement au repos (Thales/Vormetric) et en transit, durées de conservation alignées sur les obligations comptables, suppression à échéance.", "M10"),
    ("Partiel", "Head of Engineering", "Scans de vulnérabilités trimestriels (ASV), patch management avec SLA, anti-malware sur les postes accédant aux consoles de paiement.", "M10"),
    ("Partiel", "Head of Engineering + Steward Payments", "Accès aux consoles et données de paiement limité aux rôles nommés, MFA obligatoire, revue des accès trimestrielle, journalisation nominative.", "M10"),
    ("Partiel", "Head of Engineering", "Journalisation centralisée dans Splunk, alertes sur accès anormaux, pentest annuel du parcours de paiement, revue des logs par le steward.", "M10 (setup), annuel"),
    ("Partiel", "CDO + DPO + Security", "Politique de sécurité de l'information annexée à la Data Governance Policy, communiquée à tout le personnel (e-learning obligatoire, > 90 % formés), revue annuelle.", "M2 (publication), M12 (revue)"),
]

STATUS_FILL = {
    "Oui": "C6EFCE",
    "Partiel": "FFEB9C",
    "Non": "FFC7CE",
}


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb.active
    ws.title = "Checklist"
    # colonnes supplémentaires : Owner, Échéance
    ws.cell(row=1, column=5, value="Owner")
    ws.cell(row=1, column=6, value="Échéance (plan)")
    for r, (status, owner, plan, due) in enumerate(ROWS, start=2):
        ws.cell(row=r, column=3, value=status)
        ws.cell(row=r, column=4, value=plan)
        ws.cell(row=r, column=5, value=owner)
        ws.cell(row=r, column=6, value=due)
        key = status.split(" ")[0]
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=STATUS_FILL.get(key, "FFFFFF"))
    # mise en forme
    widths = {1: 34, 2: 48, 3: 18, 4: 80, 5: 30, 6: 22}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1DB954")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=1 + len(ROWS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    # feuille de synthèse
    syn = wb.create_sheet("Synthèse")
    syn["A1"] = "Spotify — Compliance Checklist (GDPR · CCPA/CPRA · PCI-DSS)"
    syn["A1"].font = Font(bold=True, size=13)
    syn["A2"] = "Statut = état actuel estimé à l'issue du Data Maturity Assessment (hypothèses, sans accès aux données internes). Action Plan = mesures de la Data Governance Policy et du plan d'implémentation. Échéance = mois du plan (M1-M18)."
    syn["A2"].alignment = Alignment(wrap_text=True)
    syn.column_dimensions["A"].width = 110
    counts = {}
    for status, *_ in ROWS:
        counts[status.split(" ")[0]] = counts.get(status.split(" ")[0], 0) + 1
    syn["A4"] = "Répartition des 14 exigences"
    syn["A4"].font = Font(bold=True)
    r = 5
    for k in ("Oui", "Partiel", "Non"):
        syn.cell(row=r, column=1, value=f"{k} : {counts.get(k, 0)}")
        r += 1
    syn.cell(row=r + 1, column=1, value="Préparé par : Emeline ROBLOT — Data Governance Specialist — Septembre 2026")

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("checklist ->", OUT)


if __name__ == "__main__":
    main()
